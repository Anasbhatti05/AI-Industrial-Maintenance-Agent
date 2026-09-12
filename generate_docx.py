from docx import Document
import os
import zipfile

output_path = r"C:\Users\DELL\OneDrive\Desktop\AI-Industrial-Maintenance-Agent\AI_Industrial_Maintenance_Agent_Document.docx"

# Create document with clean content

doc = Document()
doc.add_heading('AI Industrial Maintenance Agent', 0)

doc.add_heading('1. Project Overview', level=1)
doc.add_paragraph('The AI Industrial Maintenance Agent is an intelligent assistant designed to help industrial teams diagnose equipment issues, retrieve maintenance knowledge, and recommend corrective actions using AI, retrieval-augmented generation (RAG), and structured maintenance data.')

doc.add_heading('2. Problem Statement', level=1)
doc.add_paragraph('Industrial plants and factories rely on equipment manuals, maintenance logs, SOPs, troubleshooting guides, and historical issue records. These documents are often distributed across systems, making it difficult for maintenance teams to find the right information quickly. This leads to longer downtime and higher operational cost.')

doc.add_heading('3. Core Idea', level=1)
doc.add_paragraph('The system receives a maintenance problem such as overheating, vibration, pressure issues, or abnormal noise. It then retrieves relevant documentation and historical records, evaluates probable causes, and provides a structured recommendation for technicians.')

doc.add_heading('4. Proposed Workflow', level=1)
for step in [
    'Step 1: Technician enters equipment issue and machine details.',
    'Step 2: System retrieves manuals, SOPs, and historical maintenance data.',
    'Step 3: The AI agent searches for similar past cases and likely root causes.',
    'Step 4: It recommends troubleshooting steps, safety checks, and spare parts.',
    'Step 5: It generates a technician-friendly summary and maintenance report.'
]:
    doc.add_paragraph(step, style='List Bullet')

doc.add_heading('5. Why It Matters', level=1)
doc.add_paragraph('This project solves a valuable real-world problem: technicians spend too much time finding information manually. The AI agent helps them work faster, reduce downtime, and improve maintenance quality.')

doc.add_heading('6. Technology Stack', level=1)
for item in [
    'Frontend: Streamlit or a simple web app',
    'Backend: FastAPI',
    'LLM: OpenAI / Claude / Gemini',
    'RAG: FAISS, Chroma, or Pinecone',
    'Document ingestion: PDF parsing and chunking',
    'Deployment: Render, Railway, or Docker'
]:
    doc.add_paragraph(item, style='List Bullet')

doc.add_heading('7. MVP Scope', level=1)
for item in [
    'Upload equipment manuals and SOPs',
    'Enter a machine fault description',
    'Retrieve relevant documents',
    'Show probable root cause and troubleshooting steps',
    'Generate a structured summary for maintenance teams'
]:
    doc.add_paragraph(item, style='List Bullet')

doc.add_heading('8. Final Summary', level=1)
doc.add_paragraph('The AI Industrial Maintenance Agent is a practical and valuable solution for modern industrial operations. It combines AI, retrieval, and domain knowledge to create a maintenance assistant that helps teams diagnose issues faster and reduce downtime.')

doc.save(output_path)

# Validate file is a real zip-based DOCX
with zipfile.ZipFile(output_path) as z:
    entries = z.namelist()
    print('VALID_DOCX=True')
    print('ENTRY_COUNT=' + str(len(entries)))
    print('FIRST_ENTRIES=' + ', '.join(entries[:5]))

print('FILE_SAVED=' + output_path)
print('FILE_SIZE=' + str(os.path.getsize(output_path)))
